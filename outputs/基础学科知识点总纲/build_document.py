from pathlib import Path
import re, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
SOURCE = BASE / '知识点总纲.md'
OUT = BASE / '小学初高中基础学科知识点总纲.docx'
FONT = 'Hiragino Sans GB'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.68)
sec.bottom_margin = Inches(.66)
sec.left_margin = sec.right_margin = Inches(.72)
sec.header_distance = sec.footer_distance = Inches(.28)
WIDTH = 8.5 - 1.44

def set_font(obj, size=None, bold=None, color='000000'):
    obj.font.name = FONT
    if size is not None: obj.font.size = Pt(size)
    if bold is not None: obj.font.bold = bold
    obj.font.color.rgb = RGBColor.from_string(color)
    el = obj.element if hasattr(obj, 'element') else obj._element
    rp = el.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.insert(0, rf)
    for a in ('ascii','hAnsi','eastAsia','cs'): rf.set(qn('w:'+a), FONT)
    for a in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        rf.attrib.pop(qn('w:'+a), None)
    color_el = rp.find(qn('w:color'))
    if color_el is not None:
        for a in ('themeColor','themeTint','themeShade'): color_el.attrib.pop(qn('w:'+a), None)

for st in doc.styles:
    if st.type == 1:
        set_font(st, 11)
        if st.element.pPr is not None:
            for x in list(st.element.pPr):
                if x.tag in (qn('w:pBdr'), qn('w:shd')): st.element.pPr.remove(x)

style_specs = {
    'Normal':(11,False,0,5), 'Body Text':(11,False,0,5),
    'Title':(24,True,0,12), 'Subtitle':(11,False,0,7),
    'Heading 1':(18,True,14,10), 'Heading 2':(15,True,16,8),
    'Heading 3':(12,True,10,6), 'List Bullet':(11,False,0,5),
    'Caption':(9,False,3,5),
}
for name,(size,bold,before,after) in style_specs.items():
    st=doc.styles[name]; set_font(st,size,bold)
    pf=st.paragraph_format
    pf.space_before=Pt(before); pf.space_after=Pt(after)
    pf.line_spacing=Pt(17 if size<=11 else size*1.35)
    pf.widow_control=True
    pf.keep_with_next=name.startswith('Heading')
doc.styles['List Bullet'].paragraph_format.left_indent=Inches(.15)
doc.styles['List Bullet'].paragraph_format.first_line_indent=Inches(-.15)

def field(p, code, initial=''):
    r=OxmlElement('w:r'); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); r.append(begin); p._p.append(r)
    r=OxmlElement('w:r'); t=OxmlElement('w:instrText'); t.set(qn('xml:space'),'preserve'); t.text=' '+code+' '; r.append(t); p._p.append(r)
    r=OxmlElement('w:r'); sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); r.append(sep); p._p.append(r)
    if initial: p.add_run(initial)
    r=OxmlElement('w:r'); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r.append(end); p._p.append(r)

hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
set_font(hp.add_run('小学·初中·高中  /  基础学科知识点总纲'),8.5)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run('知识框架与复习自检  ·  '),8.5)
field(fp,'PAGE','1')

def inline(p, text):
    for i,s in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        if s:
            r=p.add_run(s)
            if i%2: r.bold=True

def link(p,label,url=None,anchor=None,size=10.5):
    h=OxmlElement('w:hyperlink')
    if url:
        rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
        h.set(qn('r:id'),rid)
    if anchor: h.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts')
    for k in ('ascii','hAnsi','eastAsia'): rf.set(qn('w:'+k),FONT)
    rp.append(rf)
    co=OxmlElement('w:color'); co.set(qn('w:val'),'000000' if anchor else '175C82'); rp.append(co)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(size*2))); rp.append(sz)
    if url:
        u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u)
    r.append(rp); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

bookmark_id=0
def bookmark(p,name):
    global bookmark_id
    bookmark_id+=1
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(bookmark_id)); start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(bookmark_id))
    p._p.insert(0,start); p._p.append(end)

def add_table(lines):
    rows=[[c.strip() for c in x.strip().strip('|').split('|')] for x in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[-: ]+',c or '-') for c in r)]
    n=len(rows[0]); widths=([1.05,1.7,1.8,2.51] if n==4 else [1.0,2.86,3.20])
    t=doc.add_table(rows=0,cols=n); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for grid,w in zip(t._tbl.tblGrid.gridCol_lst,widths): grid.set(qn('w:w'),str(int(w*1440)))
    for ri,rowdata in enumerate(rows):
        row=t.add_row()
        trPr=row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if ri==0: trPr.append(OxmlElement('w:tblHeader'))
        for ci,(c,txt) in enumerate(zip(row.cells,rowdata)):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side,val in [('top',80),('bottom',80),('start',100),('end',100)]:
                z=OxmlElement('w:'+side); z.set(qn('w:w'),str(val)); z.set(qn('w:type'),'dxa'); mar.append(z)
            tcpr.append(mar)
            bord=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                z=OxmlElement('w:'+side); z.set(qn('w:val'),'single'); z.set(qn('w:sz'),'5'); z.set(qn('w:color'),'D9D9D9'); bord.append(z)
            tcpr.append(bord)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EEF2' if ri==0 else 'FFFFFF'); tcpr.append(shade)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=Pt(14)
            if ri==0 or ci==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(txt),10,ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=SOURCE.read_text().splitlines()
chapters=[]
stage=''
for ln in lines:
    if ln.startswith('## ') and any(x in ln for x in ['第一部分','第二部分','第三部分']): stage=ln[3:]
    if re.match(r'### \d\d',ln): chapters.append((stage,ln[4:],f'ch_{len(chapters)+1:02d}'))
page_map_path=BASE/'page_map.json'
page_map=json.loads(page_map_path.read_text()) if page_map_path.exists() else {}

def add_contents():
    doc.add_page_break()
    doc.add_heading('学科目录',level=1)
    p=doc.add_paragraph('点击学科名称可跳转。小学 5 科、初中 9 科、高中 9 科；同一学科按学段递进。')
    p.paragraph_format.space_after=Pt(8)
    previous=''
    for stage,title,anchor in chapters:
        if stage!=previous:
            p=doc.add_paragraph()
            p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.keep_with_next=True
            set_font(p.add_run(stage),11,True)
            previous=stage
        p=doc.add_paragraph()
        p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH-.06), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        link(p,title,anchor=anchor,size=10.5)
        p.add_run('\t')
        field(p,'PAGEREF '+anchor+' \\h',str(page_map.get(anchor,'')))

i=0; body_started=False; chapter_index=0
while i<len(lines):
    ln=lines[i].strip(); i+=1
    if not ln: continue
    if ln.startswith('|'):
        tablelines=[ln]
        while i<len(lines) and lines[i].strip().startswith('|'):
            tablelines.append(lines[i].strip()); i+=1
        add_table(tablelines); continue
    if ln.startswith('# '):
        title=ln[2:].replace('基础学科知识点总纲','\n基础学科知识点总纲')
        doc.add_paragraph(title,'Title'); continue
    if ln.startswith('## '):
        title=ln[3:]
        if title.startswith('第一部分'):
            add_contents(); body_started=True
        p=doc.add_heading(title,level=1)
        if body_started: p.paragraph_format.page_break_before=True
        continue
    if ln.startswith('### '):
        title=ln[4:]; p=doc.add_heading(title,level=2)
        if re.match(r'\d\d',title):
            chapter_index+=1; bookmark(p,f'ch_{chapter_index:02d}')
            if title.startswith('23'): p.paragraph_format.page_break_before=True
        continue
    if ln.startswith('#### '):
        doc.add_heading(ln[5:],level=3); continue
    match=re.fullmatch(r'\[([^\]]+)\]\((.+)\)',ln)
    if match:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(7)
        link(p,match[1],url=match[2]); continue
    if ln.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); inline(p,ln[2:])
        p.paragraph_format.keep_together=True
        continue
    p=doc.add_paragraph(style='Body Text'); inline(p,ln)
    if ln.startswith('参考体系'):
        for r in p.runs: set_font(r,9.5)
        p.paragraph_format.space_after=Pt(8); p.paragraph_format.keep_with_next=True
    elif ln.startswith('整理日期') or ln.startswith('统编版与'):
        for r in p.runs: set_font(r,10)
    elif ln.startswith('［'):
        p.paragraph_format.keep_with_next=True

# Keep all headings black and free of legacy theme borders.
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name=='Title':
        for r in p.runs: set_font(r,bold=True)
        if p._p.pPr is not None:
            for x in list(p._p.pPr):
                if x.tag==qn('w:pBdr'): p._p.pPr.remove(x)
    p.paragraph_format.widow_control=True

settings=doc.settings.element
up=OxmlElement('w:updateFields'); up.set(qn('w:val'),'true'); settings.append(up)
doc.core_properties.title='小学初高中基础学科知识点总纲'
doc.core_properties.subject='统编版、人教版与教科版通用知识框架'
doc.core_properties.author=''
doc.core_properties.keywords='小学,初中,高中,基础学科,知识点,复习'
doc.save(OUT)
print(json.dumps({'output':str(OUT),'chapters':len(chapters),'knowledge_modules':sum(x.startswith('- **') for x in lines),'characters':len(SOURCE.read_text())},ensure_ascii=False))
